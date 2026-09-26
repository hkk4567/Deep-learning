#!/usr/bin/env python3
"""
log_change.py - Ghi nhật ký thay đổi (append-only) vào file change_log.xlsx.

Sheets:
  README        : chú giải cột, quy ước
  Change_Log    : mỗi hành động/khám phá/thay đổi = 1 dòng
  Findings      : các phát hiện (lỗi, bất thường) có bằng chứng
  File_Registry : danh mục file (kích thước, SHA256, vai trò)
  Experiments   : mỗi thí nghiệm DL = 1 dòng (config, seed, metric, checkpoint)
  Summary       : thống kê bằng công thức COUNTIF (chạy recalc.py sau khi ghi xong)

Lệnh:
  init      --log LOG
  add       --log LOG --actor A --action FIX --file F --desc "..." [...]
  finding   --log LOG --category C --severity High --desc "..." [...]
  register  --log LOG --file PATH --role "..." [--log-id LOG-0001]
  experiment  --log LOG --goal "..." --model resnet18 --config PATH ...   (thêm dòng ở sheet Experiments)
  exp-update  --log LOG --exp-id EXP-0001 --status Done --val-metric 0.91 ...
  backup    --file PATH --backup-dir DIR           (in ra đường dẫn bản sao)
  batch     --log LOG --json entries.json          (list các dict cho `add`)

Ví dụ:
  python log_change.py init --log ws/04_logs/change_log.xlsx
  python log_change.py add --log ws/04_logs/change_log.xlsx --actor Claude \
      --action FIX --step "B7 Chuẩn hóa" --file preprocessing.ipynb \
      --path 01_workspace/preprocessing/preprocessing.ipynb --location "cell 54" \
      --desc "Ghi df_standard thay vì df_minmax" --reason "Lỗi copy-paste" \
      --before "df_minmax.to_excel(df_standard.xlsx)" --after "df_standard.to_excel(...)" \
      --severity High --status Done --verify "so sánh 2 xlsx: khác nhau" --backup 05_backups/...
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import shutil
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo(os.environ.get("LOG_TZ", "Asia/Ho_Chi_Minh"))
except Exception:  # pragma: no cover
    TZ = None

FONT = "Arial"
ACTIONS = ["EXPLORE", "DISCOVER", "CREATE", "MODIFY", "FIX", "DELETE",
           "RUN", "VERIFY", "DECISION", "NOTE"]
SEVERITIES = ["Info", "Low", "Medium", "High", "Critical"]
STATUSES = ["Done", "Open", "Needs-confirmation", "Reverted", "Wontfix"]
ACTOR_TYPES = ["AI", "Human", "Tool"]

# (header, key, width)
LOG_COLS = [
    ("Log_ID", "log_id", 10),
    ("Ngày (dd/mm/yyyy)", "date", 14),
    ("Giờ (hh:mm:ss)", "time", 11),
    ("Múi_giờ", "tz", 14),
    ("Session_ID", "session", 14),
    ("Người_thực_hiện", "actor", 16),
    ("Loại_actor", "actor_type", 10),
    ("Yêu_cầu_bởi", "requested_by", 16),
    ("Loại_hành_động", "action", 14),
    ("Bước_pipeline", "step", 20),
    ("Tên_file", "file", 26),
    ("Đường_dẫn", "path", 44),
    ("Vị_trí (dòng/cell/hàm)", "location", 22),
    ("Nội_dung", "desc", 60),
    ("Lý_do", "reason", 40),
    ("Trước", "before", 40),
    ("Sau", "after", 40),
    ("Mức_độ", "severity", 11),
    ("Trạng_thái", "status", 18),
    ("Cách_kiểm_chứng", "verify", 40),
    ("Liên_quan_Log_ID", "related", 16),
    ("Đường_dẫn_backup", "backup", 40),
    ("Lệnh_Công_cụ", "tool", 34),
    ("Ghi_chú", "notes", 34),
]
FIND_COLS = [
    ("Finding_ID", "finding_id", 11),
    ("Ngày", "date", 12),
    ("Mức_độ", "severity", 11),
    ("Nhóm", "category", 22),
    ("Bước_pipeline", "step", 20),
    ("Tên_file", "file", 26),
    ("Vị_trí", "location", 22),
    ("Mô_tả", "desc", 60),
    ("Bằng_chứng", "evidence", 50),
    ("Cách_sửa_đề_xuất", "fix", 44),
    ("Trạng_thái", "status", 18),
    ("Log_ID_liên_quan", "related", 16),
]
REG_COLS = [
    ("Đường_dẫn", "path", 56),
    ("Loại", "ftype", 10),
    ("Kích_thước_byte", "size", 16),
    ("SHA256", "sha256", 66),
    ("Vai_trò", "role", 40),
    ("Ngày_đăng_ký", "date", 14),
    ("Log_ID_đăng_ký", "log_id", 14),
]

EXP_STATUSES = ["Planned", "Running", "Done", "Failed", "Abandoned"]
EXP_COLS = [
    ("Exp_ID", "exp_id", 10),
    ("Ngày", "date", 12),
    ("Giờ", "time", 10),
    ("Người_thực_hiện", "actor", 16),
    ("Mục_tiêu/Giả_thuyết", "goal", 44),
    ("Exp_cha", "parent", 10),
    ("Thay_đổi_so_với_exp_cha", "change", 40),
    ("Model", "model", 20),
    ("Số_tham_số", "n_params", 14),
    ("Phiên_bản_dữ_liệu/Split", "data_version", 22),
    ("Đường_dẫn_config", "config", 34),
    ("Hash_config", "config_hash", 18),
    ("Seed", "seed", 8),
    ("Epochs", "epochs", 9),
    ("Batch", "batch", 8),
    ("Optimizer", "optimizer", 14),
    ("LR", "lr", 10),
    ("Augmentation", "aug", 28),
    ("Tên_metric", "metric_name", 14),
    ("Train_metric", "train_metric", 13),
    ("Val_metric", "val_metric", 13),
    ("Test_metric (chỉ mô hình cuối)", "test_metric", 18),
    ("Thời_gian_train", "train_time", 14),
    ("Phần_cứng", "hardware", 20),
    ("Checkpoint", "ckpt", 34),
    ("Kết_luận", "conclusion", 44),
    ("Trạng_thái", "status", 12),
    ("Log_ID_liên_quan", "related", 16),
]

README_ROWS = [
    ("change_log.xlsx - Nhật ký thay đổi dự án ML", ""),
    ("Quy ước", "Chỉ THÊM dòng (append-only). Không sửa/xóa dòng cũ; muốn đính chính thì thêm dòng mới và điền Liên_quan_Log_ID."),
    ("Loại_hành_động", "EXPLORE đọc/khám phá | DISCOVER phát hiện mới | CREATE tạo file | MODIFY sửa | FIX sửa lỗi | DELETE xóa | RUN chạy code/lệnh | VERIFY kiểm chứng | DECISION quyết định (thường do người dùng) | NOTE ghi chú"),
    ("Người_thực_hiện", "Ai thực hiện hành động: 'Claude' hoặc tên người dùng/người khác."),
    ("Yêu_cầu_bởi", "Ai yêu cầu/quyết định hành động (vd tên người dùng)."),
    ("Mức_độ", "Info < Low < Medium < High < Critical."),
    ("Trạng_thái", "Done | Open | Needs-confirmation | Reverted | Wontfix."),
    ("Trước/Sau", "Giá trị, dòng code hoặc hash trước và sau khi thay đổi."),
    ("Múi_giờ", "Mặc định Asia/Ho_Chi_Minh (đổi bằng biến môi trường LOG_TZ)."),
    ("Experiments", "Mỗi thí nghiệm DL một dòng: Exp_ID, config + hash, seed, metric train/val, checkpoint. Test_metric chỉ điền cho mô hình cuối (test set dùng đúng một lần)."),
    ("Summary", "Thống kê bằng công thức; chạy recalc.py sau khi ghi xong để có giá trị cache."),
]

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SEV_FILL = {"High": "FCE4D6", "Critical": "F8CBAD"}


def now(at=None):
    if at:
        return dt.datetime.strptime(at, "%Y-%m-%d %H:%M:%S")
    n = dt.datetime.now(TZ) if TZ else dt.datetime.now()
    return n.replace(tzinfo=None)


def tzname():
    return os.environ.get("LOG_TZ", "Asia/Ho_Chi_Minh") if TZ else "local"


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def style_header(ws, cols):
    for i, (h, _, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"


def add_list_validation(ws, cols, key, values, last_row=5000):
    idx = [k for _, k, _ in cols].index(key) + 1
    col = get_column_letter(idx)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last_row}")


def init_workbook(path):
    if os.path.exists(path):
        return
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for r, (a, b) in enumerate(README_ROWS, 1):
        ws.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=True, size=12 if r == 1 else 10)
        c = ws.cell(row=r, column=2, value=b)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Change_Log")
    style_header(ws, LOG_COLS)
    add_list_validation(ws, LOG_COLS, "action", ACTIONS)
    add_list_validation(ws, LOG_COLS, "severity", SEVERITIES)
    add_list_validation(ws, LOG_COLS, "status", STATUSES)
    add_list_validation(ws, LOG_COLS, "actor_type", ACTOR_TYPES)

    ws = wb.create_sheet("Findings")
    style_header(ws, FIND_COLS)
    add_list_validation(ws, FIND_COLS, "severity", SEVERITIES)
    add_list_validation(ws, FIND_COLS, "status", STATUSES)

    ws = wb.create_sheet("File_Registry")
    style_header(ws, REG_COLS)

    ws = wb.create_sheet("Experiments")
    style_header(ws, EXP_COLS)
    add_list_validation(ws, EXP_COLS, "status", EXP_STATUSES)

    ws = wb.create_sheet("Summary")
    ws["A1"] = "Thống kê nhật ký (công thức - chạy recalc.py để cập nhật giá trị cache)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws["A3"], ws["B3"] = "Tổng số dòng log", "=COUNTA(Change_Log!A2:A5000)"
    ws["A4"], ws["B4"] = "Tổng số finding", "=COUNTA(Findings!A2:A5000)"
    ws["A5"], ws["B5"] = "Finding còn Open", '=COUNTIF(Findings!K2:K5000,"Open")'
    ws["D3"], ws["E3"] = "Tổng số thí nghiệm", "=COUNTA(Experiments!A2:A5000)"
    ws["D4"], ws["E4"] = "Thí nghiệm Done", '=COUNTIF(Experiments!AA2:AA5000,"Done")'
    ws["D5"], ws["E5"] = "Thí nghiệm Failed", '=COUNTIF(Experiments!AA2:AA5000,"Failed")'
    ws.column_dimensions["D"].width = 24
    r = 7
    ws.cell(row=r, column=1, value="Theo loại hành động").font = Font(name=FONT, bold=True)
    for a in ACTIONS:
        r += 1
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=f'=COUNTIF(Change_Log!I2:I5000,A{r})')
    r += 2
    ws.cell(row=r, column=1, value="Theo trạng thái").font = Font(name=FONT, bold=True)
    for s in STATUSES:
        r += 1
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value=f'=COUNTIF(Change_Log!S2:S5000,A{r})')
    r += 2
    ws.cell(row=r, column=1, value="Theo mức độ").font = Font(name=FONT, bold=True)
    for s in SEVERITIES:
        r += 1
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value=f'=COUNTIF(Change_Log!R2:R5000,A{r})')
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != FONT:
                c.font = Font(name=FONT, size=10, bold=c.font.bold)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    wb.save(path)


def next_id(ws, prefix):
    n = 0
    for (v,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if v and str(v).startswith(prefix + "-"):
            try:
                n = max(n, int(str(v).split("-")[1]))
            except ValueError:
                pass
    return f"{prefix}-{n + 1:04d}"


def append_row(ws, cols, rec):
    row = [rec.get(k, "") for _, k, _ in cols]
    ws.append(row)
    r = ws.max_row
    for i, (_, k, _) in enumerate(cols, 1):
        c = ws.cell(row=r, column=i)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if k == "date" and isinstance(rec.get(k), dt.date):
            c.number_format = "dd/mm/yyyy"
        if k == "time" and isinstance(rec.get(k), dt.time):
            c.number_format = "hh:mm:ss"
    sev = rec.get("severity")
    if sev in SEV_FILL:
        i = [k for _, k, _ in cols].index("severity") + 1
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=SEV_FILL[sev])


def add_entry(wb, a):
    ws = wb["Change_Log"]
    t = now(a.get("at"))
    rec = {
        "log_id": next_id(ws, "LOG"),
        "date": t.date(), "time": t.time().replace(microsecond=0), "tz": tzname(),
        "session": a.get("session") or f"S-{t:%Y%m%d}",
        "actor": a.get("actor") or "Claude",
        "actor_type": a.get("actor_type") or ("AI" if (a.get("actor") or "Claude") == "Claude" else "Human"),
        "requested_by": a.get("requested_by", ""),
        "action": (a.get("action") or "NOTE").upper(),
        "step": a.get("step", ""), "file": a.get("file", ""), "path": a.get("path", ""),
        "location": a.get("location", ""), "desc": a.get("desc", ""), "reason": a.get("reason", ""),
        "before": a.get("before", ""), "after": a.get("after", ""),
        "severity": a.get("severity") or "Info", "status": a.get("status") or "Done",
        "verify": a.get("verify", ""), "related": a.get("related", ""),
        "backup": a.get("backup", ""), "tool": a.get("tool", ""), "notes": a.get("notes", ""),
    }
    if rec["action"] not in ACTIONS:
        sys.exit(f"action không hợp lệ: {rec['action']} (chọn: {', '.join(ACTIONS)})")
    if rec["severity"] not in SEVERITIES:
        sys.exit(f"severity không hợp lệ: {rec['severity']}")
    if rec["status"] not in STATUSES:
        sys.exit(f"status không hợp lệ: {rec['status']}")
    if not rec["desc"]:
        sys.exit("--desc (Nội_dung) là bắt buộc")
    append_row(ws, LOG_COLS, rec)
    return rec["log_id"]


def add_finding(wb, a):
    ws = wb["Findings"]
    t = now(a.get("at"))
    rec = {
        "finding_id": next_id(ws, "FND"), "date": t.date(),
        "severity": a.get("severity") or "Medium", "category": a.get("category", ""),
        "step": a.get("step", ""), "file": a.get("file", ""), "location": a.get("location", ""),
        "desc": a.get("desc", ""), "evidence": a.get("evidence", ""), "fix": a.get("fix", ""),
        "status": a.get("status") or "Open", "related": a.get("related", ""),
    }
    if not rec["desc"]:
        sys.exit("--desc là bắt buộc")
    append_row(ws, FIND_COLS, rec)
    return rec["finding_id"]


def register_file(wb, path, role="", log_id=""):
    ws = wb["File_Registry"]
    t = now()
    rec = {
        "path": path, "ftype": os.path.splitext(path)[1].lstrip(".").lower(),
        "size": os.path.getsize(path), "sha256": sha256(path), "role": role,
        "date": t.date(), "log_id": log_id,
    }
    for row in ws.iter_rows(min_row=2):
        if row[0].value == path:
            row[2].value, row[3].value = rec["size"], rec["sha256"]
            if role:
                row[4].value = role
            return "updated"
    append_row(ws, REG_COLS, rec)
    return "added"


def ensure_sheets(wb):
    """Workbook cũ (tạo bởi skill khác) có thể thiếu sheet Experiments."""
    if "Experiments" not in wb.sheetnames:
        idx = wb.sheetnames.index("Summary") if "Summary" in wb.sheetnames else None
        ws = wb.create_sheet("Experiments", index=idx)
        style_header(ws, EXP_COLS)
        add_list_validation(ws, EXP_COLS, "status", EXP_STATUSES)


NUMERIC_EXP = {"n_params", "seed", "epochs", "batch", "lr", "train_metric", "val_metric", "test_metric"}


def _coerce(k, v):
    """Chuyển chuỗi số thành số để sheet Experiments sort/vẽ biểu đồ được."""
    if k in NUMERIC_EXP and isinstance(v, str):
        try:
            f = float(v)
            return int(f) if f.is_integer() and "." not in v and "e" not in v.lower() else f
        except ValueError:
            return v
    return v


def _col_idx(cols, key):
    return [k for _, k, _ in cols].index(key) + 1


def _test_metric_users(ws):
    ti, ii = _col_idx(EXP_COLS, "test_metric"), _col_idx(EXP_COLS, "exp_id")
    return [r[ii - 1].value for r in ws.iter_rows(min_row=2) if r[ti - 1].value not in (None, "")]


def add_experiment(wb, a):
    ensure_sheets(wb)
    ws = wb["Experiments"]
    t = now(a.get("at"))
    rec = {"exp_id": next_id(ws, "EXP"), "date": t.date(), "time": t.time().replace(microsecond=0),
           "actor": a.get("actor") or "Claude", "status": a.get("status") or "Planned"}
    for _, k, _ in EXP_COLS:
        if k not in rec and a.get(k) is not None:
            rec[k] = _coerce(k, a[k])
    if rec["status"] not in EXP_STATUSES:
        sys.exit(f"status không hợp lệ: {rec['status']} (chọn: {', '.join(EXP_STATUSES)})")
    if not rec.get("goal"):
        sys.exit("--goal (Mục_tiêu/Giả_thuyết) là bắt buộc")
    if rec.get("test_metric") not in (None, "") and _test_metric_users(ws):
        print(f"CẢNH BÁO: test set đã được dùng bởi {_test_metric_users(ws)}; dùng lần nữa sẽ làm test hết khách quan.", file=sys.stderr)
    append_row(ws, EXP_COLS, rec)
    return rec["exp_id"]


def update_experiment(wb, a):
    ensure_sheets(wb)
    ws = wb["Experiments"]
    ii = _col_idx(EXP_COLS, "exp_id")
    for row in ws.iter_rows(min_row=2):
        if row[ii - 1].value == a["exp_id"]:
            ti = _col_idx(EXP_COLS, "test_metric") - 1
            if a.get("test_metric") not in (None, "") and row[ti].value in (None, ""):
                others = [x for x in _test_metric_users(ws) if x != a["exp_id"]]
                if others:
                    print(f"CẢNH BÁO: test set đã được dùng bởi {others}.", file=sys.stderr)
            if a.get("status") and a["status"] not in EXP_STATUSES:
                sys.exit(f"status không hợp lệ: {a['status']}")
            changed = []
            for _, k, _ in EXP_COLS:
                if k in ("exp_id", "date", "time") or a.get(k) is None:
                    continue
                row[_col_idx(EXP_COLS, k) - 1].value = _coerce(k, a[k])
                changed.append(k)
            return f"{a['exp_id']} updated: {', '.join(changed) or '(không có trường nào)'}"
    sys.exit(f"không tìm thấy {a['exp_id']}")


def do_backup(file, backup_dir):
    os.makedirs(backup_dir, exist_ok=True)
    t = now()
    base = os.path.basename(file)
    dst = os.path.join(backup_dir, f"{t:%Y%m%d_%H%M%S}__{base}")
    shutil.copy2(file, dst)
    return dst


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    def common(sp):
        sp.add_argument("--log", required=True)
        sp.add_argument("--at", help="Ghi đè thời điểm: 'YYYY-MM-DD HH:MM:SS'")

    sp = sub.add_parser("init"); common(sp)

    sp = sub.add_parser("add"); common(sp)
    for k in ["actor", "actor_type", "requested_by", "session", "action", "step", "file", "path",
              "location", "desc", "reason", "before", "after", "severity", "status", "verify",
              "related", "backup", "tool", "notes"]:
        sp.add_argument(f"--{k.replace('_', '-')}", dest=k)

    sp = sub.add_parser("finding"); common(sp)
    for k in ["severity", "category", "step", "file", "location", "desc", "evidence", "fix",
              "status", "related"]:
        sp.add_argument(f"--{k}", dest=k)

    exp_keys = [k for _, k, _ in EXP_COLS if k not in ("exp_id", "date", "time")]
    sp = sub.add_parser("experiment"); common(sp)
    for k in exp_keys:
        sp.add_argument(f"--{k.replace('_', '-')}", dest=k)
    sp = sub.add_parser("exp-update"); common(sp)
    sp.add_argument("--exp-id", required=True, dest="exp_id")
    for k in exp_keys:
        sp.add_argument(f"--{k.replace('_', '-')}", dest=k)

    sp = sub.add_parser("register"); common(sp)
    sp.add_argument("--file", required=True)
    sp.add_argument("--role", default="")
    sp.add_argument("--log-id", default="")

    sp = sub.add_parser("backup")
    sp.add_argument("--file", required=True)
    sp.add_argument("--backup-dir", required=True)

    sp = sub.add_parser("batch"); common(sp)
    sp.add_argument("--json", required=True)

    a = p.parse_args()

    if a.cmd == "backup":
        print(do_backup(a.file, a.backup_dir))
        return
    if a.cmd == "init":
        init_workbook(a.log)
        print(f"OK: {a.log}")
        return

    init_workbook(a.log)
    wb = load_workbook(a.log)
    ensure_sheets(wb)
    d = {k: v for k, v in vars(a).items() if v is not None}
    if a.cmd == "add":
        print(add_entry(wb, d))
    elif a.cmd == "finding":
        print(add_finding(wb, d))
    elif a.cmd == "experiment":
        print(add_experiment(wb, d))
    elif a.cmd == "exp-update":
        print(update_experiment(wb, d))
    elif a.cmd == "register":
        print(register_file(wb, a.file, a.role, a.log_id))
    elif a.cmd == "batch":
        with open(a.json, encoding="utf-8") as f:
            entries = json.load(f)
        ids = [add_entry(wb, {**e, "at": e.get("at") or a.at}) for e in entries]
        print(", ".join(ids))
    wb.save(a.log)


if __name__ == "__main__":
    main()
